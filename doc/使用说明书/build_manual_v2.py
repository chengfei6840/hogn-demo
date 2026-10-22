"""从 V2.0 Markdown 生成横向 A4 docx。样式沿用 build_manual.py。"""
from pathlib import Path
import re, sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
NAME = '红光金具智能报价系统_用户使用说明书_2026-10-22_V2.0'
MD = ROOT / f'{NAME}.md'
DATE = '2026年10月22日'
HEADER = '红光金具智能报价系统  用户使用说明书'

NAV_GROUPS = [
    ('开始体验', '第1至5节', '认识界面  创建项目  工作表解析  进入报价包  批量选择'),
    ('资料与需求', '第6至12节', '本包演示资料  解析与重新解析  销售确认'),
    ('技术核对', '第13至23节', '串型候选  子件与来源  参考件  跨包指派  总监确认  BOM'),
    ('成本与审批', '第24至31节', '子件成本  参考成本  演示口径  毛利率  核价提交  审批签发'),
    ('对外报价', '第32至35节', '下载外发  调整与新版本  成交结果  历史库'),
    ('管理与查阅', '第36至42节', '用户启停  角色权限  流程模板  资料治理  规则与后续范围'),
    ('操作问题', '第43节', '按钮状态  缺项检查  本地保存与重置'),
]


def set_cjk(style, name='微软雅黑'):
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)


def fmt_table(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for rownum, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if rownum == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = widths[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                b = OxmlElement('w:' + edge)
                b.set(qn('w:val'), 'single')
                b.set(qn('w:sz'), '5')
                b.set(qn('w:color'), 'D9D9D9')
                borders.append(b)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                m = OxmlElement('w:' + edge)
                m.set(qn('w:w'), '120')
                m.set(qn('w:type'), 'dxa')
                margins.append(m)
            tcpr.append(margins)
            if rownum == 0:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), 'F2F6F9')
                tcpr.append(shd)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(9)


def build():
    text = MD.read_text(encoding='utf-8')
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.3)
    sec.left_margin, sec.right_margin = Cm(1.5), Cm(1.5)
    sec.header_distance, sec.footer_distance = Cm(.5), Cm(.5)

    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
        st = doc.styles[name]
        st.font.name = 'Arial'
        set_cjk(st)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_after = Pt(5)
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.13
    doc.styles['Heading 1'].font.size = Pt(17)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
    doc.styles['Heading 2'].font.size = Pt(13)
    doc.styles['Title'].font.size = Pt(28)
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].paragraph_format.space_before = Pt(3)
    doc.styles['Caption'].paragraph_format.space_after = Pt(0)
    for st in doc.styles:
        for bd in list(st.element.iter(qn('w:pBdr'))):
            bd.getparent().remove(bd)
        if st.name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption',
                       'Title Char', 'Subtitle Char', 'Heading 1 Char', 'Heading 2 Char']:
            st.font.name = 'Arial'
            st.font.bold = False
            st.font.italic = False
            fonts = st.element.get_or_add_rPr().rFonts
            for k in list(fonts.attrib):
                if k.lower().endswith('theme'):
                    del fonts.attrib[k]
            fonts.set(qn('w:eastAsia'), '黑体')
            fonts.set(qn('w:cs'), 'Arial')
            bcs = st.element.get_or_add_rPr().find(qn('w:bCs'))
            if bcs is None:
                bcs = OxmlElement('w:bCs')
                st.element.get_or_add_rPr().append(bcs)
            bcs.set(qn('w:val'), '0')

    hp = sec.header.paragraphs[0]
    hp.text = HEADER
    hp.style = 'Normal'
    hp.runs[0].font.size = Pt(8)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.add_run(f'{DATE}  ·  第 ')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    fp._p.append(fld)
    fp.add_run(' 页')
    for r in fp.runs:
        r.font.size = Pt(8)

    # 封面
    doc.add_paragraph('红光金具智能报价系统', 'Title')
    doc.add_paragraph('用户使用说明书', 'Title')
    doc.add_paragraph(f'原型演示版  V2.0  {DATE}', 'Subtitle')
    doc.add_paragraph('本说明帮助您在原型中自行完成项目创建、物料清单工作表解析、报价包批量处理、需求及技术确认、成本核算、审批签发、销售外发和成交归档，并查阅相关管理功能。按步骤切换对应角色和人员，即可体验完整报价流程。')
    doc.add_paragraph('首次体验建议使用内置演示文档并选择包1，先只勾选“一般线路金具与防舞金具”工作表，取消两个测试工作表。截图取自当前线上原型的实际页面，放在相应操作内容下方；不同浏览器窗口、已有数据或权限设置可能使画面略有差异。')
    doc.add_paragraph('当前为本地保存的交互原型。资料解析、库价、技术推荐及服务集成存在演示数据；说明书中的技术输入值和报价金额仅用于操作体验，正式业务以经确认资料及审批结果为准。')
    doc.add_paragraph('准备事项', 'Heading 1')
    doc.add_paragraph('使用提供给您的访问地址，在同一个浏览器连续操作，数据保存在该浏览器本地。顶部右侧可切换演示角色及人员，无须依靠讲解人员代为操作；遇到按钮不可用时请先查看第43节。')

    doc.add_page_break()
    doc.add_paragraph('阅读导航', 'Heading 1')
    table = doc.add_table(rows=1, cols=3)
    for c, t in zip(table.rows[0].cells, ['阅读目的', '对应章节', '内容']):
        c.text = t
    for row in NAV_GROUPS:
        for c, t in zip(table.add_row().cells, row):
            c.text = t
    fmt_table(table, [Cm(4), Cm(4), Cm(17)])
    doc.add_paragraph('')
    doc.add_paragraph('本版相较 V1.5 的主要变化：新增第3节工作表识别与解析、第5节跨包批量处理、第21节跨包技术指派；第4节调整了报价包列表的进入方式与分类筛选；第43节补充了四类新的按钮禁用原因。')

    # 正文从新一页开始；其后每节前各分页一次（见循环内 first_section 判断）
    doc.add_page_break()

    # 正文
    lines = text.split('\n')
    i = 0
    first_section = True
    in_body = False   # 首个 ## 之前是 md 的标题与版本元信息，封面已另行生成，这里全部跳过
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if not in_body:
            if line.startswith('## '):
                in_body = True
            else:
                i += 1
                continue

        if not line or line == '---':
            i += 1
            continue
        if line.startswith('# '):
            i += 1
            continue

        if line.startswith('## '):
            if not first_section:
                doc.add_page_break()
            first_section = False
            doc.add_paragraph(line[3:].strip(), 'Heading 1')
            i += 1
            continue

        if line.startswith('### '):
            doc.add_paragraph(line[4:].strip(), 'Heading 2')
            i += 1
            continue

        m = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if m:
            alt, path = m.group(1), m.group(2)
            img = ROOT / path
            if img.exists():
                doc.add_paragraph().add_run().add_picture(str(img), width=Cm(20.0))
                cap = doc.add_paragraph(alt, 'Caption')
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                print('  ⚠ 缺图:', path, file=sys.stderr)
            i += 1
            continue

        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=1, cols=len(rows[0]))
                for c, txt in zip(t.rows[0].cells, rows[0]):
                    c.text = txt
                for r in rows[1:]:
                    cells = t.add_row().cells
                    for c, txt in zip(cells, r):
                        c.text = txt
                n = len(rows[0])
                widths = [Cm(26.7 / n)] * n
                fmt_table(t, widths)
                doc.add_paragraph('')
            continue

        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(6)
        i += 1

    out = ROOT / f'{NAME}.docx'
    doc.save(out)
    section_count = text.count('\n## ')
    image_count = len(re.findall(r'!\[.*?\]\(', text))
    print('生成:', out.name)
    print('节数:', section_count)
    print('图片引用:', image_count)


if __name__ == '__main__':
    build()
